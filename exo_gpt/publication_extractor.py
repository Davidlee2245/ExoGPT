"""
Publication Extractor using OpenAI API
Extracts biomarkers and publication metadata from PDFs and saves to Excel
"""

import os
import json
import re
from typing import Dict, List, Any, Optional
from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import PyPDF2
import docx
from openai import OpenAI

# Get OpenAI API key from environment
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")

# Cached UniProt alias map for mapping protein/marker names → gene symbols
_UNIPROT_ALIAS_MAP: Optional[Dict[str, str]] = None


def _load_uniprot_alias_map() -> Dict[str, str]:
    """
    Build a mapping from various marker names (including CD antigens)
    to canonical gene symbols using the local UniProt TSV export.
    
    Examples:
        "CD133" -> "PROM1"
        "PD-L1" (if present as alias) -> "CD274"
    """
    global _UNIPROT_ALIAS_MAP
    if _UNIPROT_ALIAS_MAP is not None:
        return _UNIPROT_ALIAS_MAP

    alias_map: Dict[str, str] = {}

    try:
        project_root = Path(__file__).resolve().parent.parent
        uniprot_path = project_root / "data" / "protein" / "uniprot" / "uniprot.tsv"
        if not uniprot_path.exists():
            _UNIPROT_ALIAS_MAP = {}
            return _UNIPROT_ALIAS_MAP

        # Only load the columns we need
        df = pd.read_csv(
            uniprot_path,
            sep="\t",
            usecols=["Gene Names", "Protein names"],
            dtype=str,
            engine="c",
        )
    except Exception:
        _UNIPROT_ALIAS_MAP = {}
        return _UNIPROT_ALIAS_MAP

    for _, row in df.iterrows():
        gene_names = str(row.get("Gene Names") or "").strip()
        prot_names = str(row.get("Protein names") or "").strip()

        if not gene_names:
            continue

        # First token in Gene Names is typically the primary gene symbol
        genes = gene_names.split()
        primary = genes[0].strip().upper()
        if not primary:
            continue

        # Map all gene name aliases to the primary symbol
        for g in genes:
            g_clean = g.strip().upper()
            if g_clean and g_clean not in alias_map:
                alias_map[g_clean] = primary

        # Extract CD antigen aliases from Protein names, e.g.
        # "Prominin-1 (Antigen AC133) (CD antigen CD133)"
        if prot_names:
            # Look for patterns like "CD133", "CD24", etc.
            for m in re.findall(r"CD\d+[A-Z]?", prot_names, flags=re.IGNORECASE):
                cd_alias = m.strip().upper()
                if cd_alias and cd_alias not in alias_map:
                    alias_map[cd_alias] = primary

    _UNIPROT_ALIAS_MAP = alias_map
    return _UNIPROT_ALIAS_MAP


def _normalize_gene_symbol_from_marker(raw_name: str) -> Dict[str, str]:
    """
    Normalize a marker/antigen name to a canonical gene symbol using UniProt.
    
    Returns dict with:
      - reported_marker: original string from the publication / OpenAI
      - gene_symbol: best-guess gene symbol (uppercased) or reported_marker if unknown
      - gene_symbol_normalized: same as gene_symbol (explicit column for downstream joins)
    """
    reported = (raw_name or "").strip()
    if not reported:
        return {
            "reported_marker": "",
            "gene_symbol": "",
            "gene_symbol_normalized": "",
        }

    alias_map = _load_uniprot_alias_map()
    key = reported.upper()

    # Direct lookup (e.g., proper HGNC symbol or CD antigen)
    gene_symbol = alias_map.get(key, key)

    return {
        "reported_marker": reported,
        "gene_symbol": gene_symbol,
        "gene_symbol_normalized": gene_symbol,
    }


def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract text from PDF file."""
    try:
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        return text
    except Exception as e:
        raise Exception(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(docx_path: str) -> str:
    """Extract text from DOCX file."""
    try:
        doc = docx.Document(docx_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text
    except Exception as e:
        raise Exception(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_file(file_path: str) -> str:
    """Extract text from file based on extension."""
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext in ['.doc', '.docx']:
        return extract_text_from_docx(file_path)
    elif ext in ['.txt']:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def extract_with_openai(text: str, filename: str) -> Dict[str, Any]:
    """Use OpenAI API to extract publication info and biomarkers."""
    if not OPENAI_API_KEY:
        raise ValueError("OPENAI_API_KEY environment variable not set")
    
    client = OpenAI(api_key=OPENAI_API_KEY)
    
    # Truncate text if too long (OpenAI has token limits)
    max_chars = 100000  # Roughly 25k tokens
    if len(text) > max_chars:
        text = text[:max_chars] + "\n[... text truncated ...]"
    
    prompt = f"""Extract the following information from this research publication:

1. Publication metadata:
   - Title
   - Journal name
   - Authors (first author and last author)
   - Publication year
   - DOI (if available)
   - PMID (if available)

2. Biomarkers mentioned in the publication:
   For each biomarker, extract:
   - Gene symbol (e.g., CD274, PDCD1)
   - UniProt ID (if mentioned)
   - Analyte type (protein, miRNA, mRNA, etc.)
   - Disease studied
   - Biofluid/sample type (e.g., plasma, serum, tissue)
   - Log2FC or fold change (if mentioned)
   - P-value (if mentioned)
   - FDR (if mentioned)
   - Method used (e.g., ELISA, qRT-PCR, mass spectrometry)

Return the data as a JSON object with this structure:
{{
  "publication": {{
    "title": "...",
    "journal": "...",
    "authors": "...",
    "year": "...",
    "doi": "...",
    "pmid": "..."
  }},
  "biomarkers": [
    {{
      "gene_symbol": "...",
      "uniprot": "...",
      "analyte_type": "protein",
      "disease": "...",
      "biofluid": "...",
      "logfc": null,
      "p_value": null,
      "fdr": null,
      "method": "..."
    }}
  ]
}}

Publication text:
{text}
"""
    
    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",  # Using gpt-4o-mini for cost efficiency
            messages=[
                {"role": "system", "content": "You are a scientific data extraction assistant. Extract publication metadata and biomarker information accurately from research papers."},
                {"role": "user", "content": prompt}
            ],
            temperature=0.1,  # Low temperature for more consistent extraction
            response_format={"type": "json_object"}
        )
        
        result_text = response.choices[0].message.content
        result = json.loads(result_text)
        
        return result
    except Exception as e:
        raise Exception(f"OpenAI API error: {str(e)}")


def check_excel_exists(excel_path: str, filename: str) -> bool:
    """Check if Excel file exists and contains data for this publication."""
    if not os.path.exists(excel_path):
        return False
    
    try:
        # Try to read Excel file
        df = pd.read_excel(excel_path, sheet_name='Publications')
        
        # Check if this filename is already in the Excel
        if 'filename' in df.columns:
            if filename in df['filename'].values:
                return True
        
        # Also check by title if filename not found
        if 'title' in df.columns:
            # We can't check title without extracting, so return False
            # This means we'll re-extract if only filename doesn't match
            pass
        
        return False
    except Exception:
        # If Excel exists but can't be read, assume it's empty/corrupted
        return False


def read_existing_data(excel_path: str, filename: str) -> Optional[Dict[str, Any]]:
    """Read existing publication and biomarker data from Excel."""
    try:
        # Read publication info
        if not os.path.exists(excel_path):
            return None
            
        pub_df = pd.read_excel(excel_path, sheet_name='Publications')
        
        if pub_df.empty or 'filename' not in pub_df.columns:
            return None
            
        pub_row = pub_df[pub_df['filename'] == filename]
        
        if pub_row.empty:
            return None
        
        pub_info = pub_row.iloc[0].to_dict()
        
        # Read biomarkers
        biomarkers = []
        try:
            biomarkers_df = pd.read_excel(excel_path, sheet_name='Biomarkers')
            if not biomarkers_df.empty and 'publication_filename' in biomarkers_df.columns:
                biomarkers = biomarkers_df[biomarkers_df['publication_filename'] == filename].to_dict('records')
        except Exception:
            # Biomarkers sheet might not exist or be empty
            pass
        
        return {
            "publication": pub_info,
            "biomarkers": biomarkers
        }
    except Exception as e:
        print(f"Warning: Could not read existing Excel data: {e}")
        return None


def save_to_excel(
    excel_path: str,
    filename: str,
    publication_info: Dict[str, Any],
    biomarkers: List[Dict[str, Any]]
):
    """Save publication info and biomarkers to Excel file."""
    
    # Add filename to publication info
    publication_info['filename'] = filename
    
    # Create or load Excel file
    if os.path.exists(excel_path):
        # Load existing workbook
        wb = openpyxl.load_workbook(excel_path)
        
        # Get or create sheets
        if 'Publications' in wb.sheetnames:
            pub_sheet = wb['Publications']
            pub_df = pd.read_excel(excel_path, sheet_name='Publications')
        else:
            pub_sheet = wb.create_sheet('Publications')
            pub_df = pd.DataFrame()
        
        if 'Biomarkers' in wb.sheetnames:
            bio_sheet = wb['Biomarkers']
            bio_df = pd.read_excel(excel_path, sheet_name='Biomarkers')
        else:
            bio_sheet = wb.create_sheet('Biomarkers')
            bio_df = pd.DataFrame()
    else:
        # Create new workbook
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        pub_sheet = wb.create_sheet('Publications')
        bio_sheet = wb.create_sheet('Biomarkers')
        pub_df = pd.DataFrame()
        bio_df = pd.DataFrame()
    
    # Extract disease information and biomarker list from biomarkers
    diseases = set()
    biofluids = set()
    biomarker_genes = []
    for bm in biomarkers:
        if bm.get('disease'):
            diseases.add(str(bm['disease']))
        if bm.get('biofluid'):
            biofluids.add(str(bm['biofluid']))
        if bm.get('gene_symbol'):
            gene = str(bm['gene_symbol']).strip()
            if gene and gene not in biomarker_genes:
                biomarker_genes.append(gene)
    
    # Prepare publication row with disease and biomarker info
    pub_row = {
        'filename': filename,
        'title': publication_info.get('title', ''),
        'journal': publication_info.get('journal', ''),
        'authors': publication_info.get('authors', ''),
        'year': publication_info.get('year', ''),
        'doi': publication_info.get('doi', ''),
        'pmid': publication_info.get('pmid', ''),
        'disease': ', '.join(sorted(diseases)) if diseases else '',
        'biofluid': ', '.join(sorted(biofluids)) if biofluids else '',
        'biomarker_count': len(biomarkers),
        'biomarkers': ', '.join(sorted(biomarker_genes)) if biomarker_genes else ''
    }
    
    # Ensure all columns exist in the DataFrame
    all_pub_columns = ['filename', 'title', 'journal', 'authors', 'year', 'doi', 'pmid', 'disease', 'biofluid', 'biomarker_count', 'biomarkers']
    for col in all_pub_columns:
        if col not in pub_df.columns:
            pub_df[col] = ''
    
    # Add publication row (update if exists, append if not)
    if not pub_df.empty and 'filename' in pub_df.columns:
        if filename in pub_df['filename'].values:
            # Update existing row
            idx = pub_df[pub_df['filename'] == filename].index[0]
            for col, val in pub_row.items():
                pub_df.at[idx, col] = val
        else:
            # Append new row - ensure all columns are present
            new_row = {col: pub_row.get(col, '') for col in all_pub_columns}
            pub_df = pd.concat([pub_df, pd.DataFrame([new_row])], ignore_index=True)
    else:
        # First publication - create with all columns
        pub_df = pd.DataFrame([pub_row])
        # Ensure all columns exist
        for col in all_pub_columns:
            if col not in pub_df.columns:
                pub_df[col] = ''
    
    # Prepare biomarker rows
    bio_rows = []
    for bm in biomarkers:
        # OpenAI returns "gene_symbol" which may actually be a protein/CD marker
        # (e.g., "CD133"). We preserve the original in reported_marker and
        # map to a canonical gene symbol using UniProt aliases.
        raw_gene = bm.get('gene_symbol', '') or ''
        gene_info = _normalize_gene_symbol_from_marker(str(raw_gene))
        bio_row = {
            'publication_filename': filename,
            # Original string extracted from the paper / OpenAI
            'reported_marker': gene_info['reported_marker'],
            # Canonical gene symbol (e.g., CD133 -> PROM1) where resolvable
            'gene_symbol': gene_info['gene_symbol'],
            # Explicit normalized gene symbol column for downstream joins / matching
            'gene_symbol_normalized': gene_info['gene_symbol_normalized'],
            'uniprot': bm.get('uniprot', ''),
            'analyte_type': bm.get('analyte_type', 'protein'),
            'disease': bm.get('disease', ''),
            'biofluid': bm.get('biofluid', ''),
            'logfc': bm.get('logfc'),
            'p_value': bm.get('p_value'),
            'fdr': bm.get('fdr'),
            'method': bm.get('method', '')
        }
        bio_rows.append(bio_row)
    
    # Remove old biomarkers for this publication and add new ones
    if not bio_df.empty and 'publication_filename' in bio_df.columns:
        bio_df = bio_df[bio_df['publication_filename'] != filename]
    
    if bio_rows:
        new_bio_df = pd.DataFrame(bio_rows)
        bio_df = pd.concat([bio_df, new_bio_df], ignore_index=True)
    
    # Write to Excel with formatting
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        pub_df.to_excel(writer, sheet_name='Publications', index=False)
        bio_df.to_excel(writer, sheet_name='Biomarkers', index=False)
    
    # Apply formatting
    wb = openpyxl.load_workbook(excel_path)
    
    # Format Publications sheet
    pub_sheet = wb['Publications']
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in pub_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Format Biomarkers sheet
    bio_sheet = wb['Biomarkers']
    for cell in bio_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Auto-adjust column widths
    for sheet in [pub_sheet, bio_sheet]:
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(excel_path)


def extract_publication_data(
    file_path: str,
    publications_dir: str,
    excel_filename: str = "publications_data.xlsx"
) -> Dict[str, Any]:
    """
    Extract publication data from file.
    
    Args:
        file_path: Path to the uploaded file
        publications_dir: Directory where publications are stored
        excel_filename: Name of Excel file to save/read from
    
    Returns:
        Dict with publication info and biomarkers
    """
    filename = os.path.basename(file_path)
    excel_path = os.path.join(publications_dir, excel_filename)
    
    # Check if data already exists in Excel
    if check_excel_exists(excel_path, filename):
        print(f"Found existing data for {filename} in Excel, skipping OpenAI extraction")
        existing_data = read_existing_data(excel_path, filename)
        if existing_data:
            return existing_data
    
    # Extract text from file
    print(f"Extracting text from {filename}...")
    text = extract_text_from_file(file_path)
    
    if not text or len(text.strip()) < 100:
        raise ValueError("Could not extract sufficient text from file")
    
    # Extract with OpenAI
    print(f"Extracting data using OpenAI API...")
    extracted_data = extract_with_openai(text, filename)
    
    # Save to Excel
    print(f"Saving to Excel: {excel_path}")
    save_to_excel(
        excel_path,
        filename,
        extracted_data.get('publication', {}),
        extracted_data.get('biomarkers', [])
    )
    
    return extracted_data

