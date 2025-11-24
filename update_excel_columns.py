#!/usr/bin/env python3
"""
Update existing Excel file to add disease, biofluid, and biomarker_count columns to Publications sheet.
"""

import os
import sys
import pandas as pd
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent
excel_path = os.path.join(PROJECT_ROOT, "data", "publications", "publications_data.xlsx")

if not os.path.exists(excel_path):
    print(f"Excel file not found: {excel_path}")
    sys.exit(1)

print(f"Reading Excel file: {excel_path}")

# Read both sheets
try:
    pub_df = pd.read_excel(excel_path, sheet_name='Publications')
    bio_df = pd.read_excel(excel_path, sheet_name='Biomarkers')
except Exception as e:
    print(f"Error reading Excel: {e}")
    sys.exit(1)

print(f"Current Publications columns: {list(pub_df.columns)}")
print(f"Current Publications rows: {len(pub_df)}")
print(f"Current Biomarkers rows: {len(bio_df)}")

# Add new columns if they don't exist
new_columns = ['disease', 'biofluid', 'biomarker_count', 'biomarkers']
for col in new_columns:
    if col not in pub_df.columns:
        pub_df[col] = ''

# For each publication, extract disease and biofluid from biomarkers
for idx, row in pub_df.iterrows():
    filename = row['filename']
    
    # Get biomarkers for this publication
    pub_biomarkers = bio_df[bio_df['publication_filename'] == filename] if 'publication_filename' in bio_df.columns else pd.DataFrame()
    
    # Extract unique diseases, biofluids, and biomarker genes
    diseases = set()
    biofluids = set()
    biomarker_genes = []
    
    if not pub_biomarkers.empty:
        if 'disease' in pub_biomarkers.columns:
            diseases = set(pub_biomarkers['disease'].dropna().astype(str).unique())
            diseases.discard('')  # Remove empty strings
        if 'biofluid' in pub_biomarkers.columns:
            biofluids = set(pub_biomarkers['biofluid'].dropna().astype(str).unique())
            biofluids.discard('')  # Remove empty strings
        if 'gene_symbol' in pub_biomarkers.columns:
            genes = pub_biomarkers['gene_symbol'].dropna().astype(str).unique()
            biomarker_genes = [g.strip() for g in genes if g.strip()]
            # Remove duplicates while preserving order
            biomarker_genes = list(dict.fromkeys(biomarker_genes))
    
    # Update publication row
    pub_df.at[idx, 'disease'] = ', '.join(sorted(diseases)) if diseases else ''
    pub_df.at[idx, 'biofluid'] = ', '.join(sorted(biofluids)) if biofluids else ''
    pub_df.at[idx, 'biomarker_count'] = len(pub_biomarkers)
    pub_df.at[idx, 'biomarkers'] = ', '.join(sorted(biomarker_genes)) if biomarker_genes else ''

# Reorder columns for better readability
column_order = ['filename', 'title', 'journal', 'authors', 'year', 'doi', 'pmid', 'disease', 'biofluid', 'biomarker_count', 'biomarkers']
# Only include columns that exist
column_order = [col for col in column_order if col in pub_df.columns]
# Add any remaining columns
remaining_cols = [col for col in pub_df.columns if col not in column_order]
pub_df = pub_df[column_order + remaining_cols]

# Save updated Excel
print(f"\nUpdating Excel file...")
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    pub_df.to_excel(writer, sheet_name='Publications', index=False)
    bio_df.to_excel(writer, sheet_name='Biomarkers', index=False)

# Apply formatting
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    
    wb = openpyxl.load_workbook(excel_path)
    
    # Format Publications sheet
    pub_sheet = wb['Publications']
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in pub_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Format Biomarkers sheet
    bio_sheet = wb['Biomarkers']
    for cell in bio_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Auto-adjust column widths
    for sheet in [pub_sheet, bio_sheet]:
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(excel_path)
    print("✓ Formatting applied")
except ImportError:
    print("⚠️  openpyxl not available for formatting, but data is saved")

print(f"\n✓ Excel file updated successfully!")
print(f"New Publications columns: {list(pub_df.columns)}")
print(f"\nSample data:")
print(pub_df[['filename', 'disease', 'biofluid', 'biomarker_count', 'biomarkers']].to_string())

